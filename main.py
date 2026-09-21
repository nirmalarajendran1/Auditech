import pandas as pd
from llm_factory import ask_ai
from openpyxl import Workbook

print("--- AuditFlow AI: Phase 1 Test ---")

# 1. Get Input
control_id = "CC6.1"
control_description = input("Enter a Control Description (or press Enter for default): ")
if not control_description:
    control_description = "User access is revoked within 24 hours of termination."

print(f"\nProcessing Control: {control_description}...")

# 2. Step A: Classify the Control
print("AI is classifying the control...")
classify_sys = "Act as an IT Auditor. Reply ONLY with 'Configuration' or 'Transactional'. No other text."
classify_prompt = f"Is this control 'Configuration' (set once) or 'Transactional' (repeated)? Control: {control_description}"

control_type = ask_ai(classify_sys, classify_prompt).strip()
# specific cleanup for Gemini's tendency to add bolding sometimes
control_type = control_type.replace("*", "").replace("\n", "").strip()

print(f"AI Classification: {control_type}")

# 3. Step B: Determine Sample Size (Human-in-the-Loop)
sample_size = 0

if "Configuration" in control_type:
    print(">> Logic: Configuration control detected. Testing 1 sample.")
    sample_size = 1

else: 
    # Fallback: If it says Transactional OR anything else, we treat it as sample-based
    print("\n[!] PAUSE: This control requires sampling.")
    pop_input = input(">> Please enter the Total Population Size (e.g., 50): ")
    risk_input = input(">> Is Risk High, Moderate, or Low? (H/M/L): ")
    
    try:
        population_size = int(pop_input)
        if risk_input.lower() == 'h':
            sample_size = min(40, population_size)
        elif risk_input.lower() == 'm':
            sample_size = min(25, population_size)
        else:
            sample_size = min(10, population_size)
    except ValueError:
        print("Invalid number. Defaulting to 5 samples.")
        sample_size = 5
        
    print(f">> Calculated Sample Size: {sample_size}")

# 4. Step C: Generate Test Procedure
print("\nGenerating Test Procedure...")
proc_sys = "Act as a SOX Auditor. Write a clear, 3-sentence test procedure starting with a verb (Inspect, Observe)."
proc_prompt = f"Write a test procedure for: {control_description}"

test_procedure = ask_ai(proc_sys, proc_prompt)

# 5. Step D: Write to Excel
print("\nSaving to Excel...")
wb = Workbook()
ws = wb.active
ws.title = control_id

# Header Info
ws['A1'] = "Control ID"
ws['B1'] = control_id
ws['A2'] = "Description"
ws['B2'] = control_description
ws['A3'] = "Type"
ws['B3'] = control_type
ws['A4'] = "Test Procedure"
ws['B4'] = test_procedure

# The Testing Table
ws['A6'] = "Sample #"
ws['B6'] = "Test Attribute 1"
ws['C6'] = "Test Attribute 2"
ws['D6'] = "Result (Pass/Fail)"

# Create blank rows
for i in range(1, sample_size + 1):
    row_num = 6 + i
    ws[f'A{row_num}'] = f"Sample {i}"

filename = "Audit_Test_Sheet.xlsx"
wb.save(filename)
print(f"SUCCESS! File saved as '{filename}'")